#!/usr/bin/env python3
"""Build docs/data.json for the GitHub Pages village-match explorer.

Reads FuzzyMatch/Lebanese_Villages_Matched_v6.xlsx and emits a compact JSON
with matched pairs, review candidates, and unmatched villages on both sides,
each tagged with a canonical district key so the page can show same-district
unmatched villages when there is no definite match.
"""
import argparse
import json
import re
import unicodedata
from collections import defaultdict
from difflib import SequenceMatcher
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parent
XLSX = REPO / "FuzzyMatch" / "Lebanese_Villages_Matched_v6.xlsx"
OUT = REPO / "docs" / "data.json"

_ap = argparse.ArgumentParser()
_ap.add_argument("--input", default=str(XLSX), help="matched workbook (.xlsx)")
_ap.add_argument("--output", default=str(OUT), help="output data.json")
_args = _ap.parse_args()
XLSX = Path(_args.input)
OUT = Path(_args.output)


def split_ar_cell(s: str):
    return [p.strip() for p in re.split(r"[,،;]", s or "") if p.strip()]


def norm_ar(s: str) -> str:
    s = s or ""
    s = re.sub(r"[\u064b-\u0652\u0670]", "", s)  # tashkeel
    s = s.replace("\u0640", "")  # tatweel
    s = re.sub(r"[أإآٱ]", "ا", s)
    s = s.replace("ى", "ي").replace("ة", "ه")
    s = re.sub(r"\s+", " ", s).strip()
    return s


def norm_lat(s: str) -> str:
    s = s or ""
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", s).strip().lower()


def clean_ar_district(s: str) -> str:
    s = s or ""
    s = re.sub(r"\s*\([^)]*\)\s*", " ", s)  # drop (عكار) style suffixes
    s = re.sub(r"^(قضاءي?|قرى|محافظة)\s+", "", s.strip())
    return norm_ar(s)


wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)

# --- translation tables -----------------------------------------------------
dist_rows = list(wb["DistrictTranslation"].iter_rows(values_only=True))[1:]
# (District latin, District Arabic, Mohafaza)
lat2ar_d, ar2lat_d = {}, {}
for latin, arabic, _moh in dist_rows:
    if not latin or not arabic:
        continue
    lat2ar_d[norm_lat(latin)] = arabic
    ar2lat_d[norm_ar(arabic)] = latin
KNOWN_AR_D = list(ar2lat_d.keys())

moh_rows = list(wb["Mohafaza Translations"].iter_rows(values_only=True))[1:]
# (Mohafaza Arabic, In English)
ar2lat_m, lat2ar_m = {}, {}
for arabic, latin in moh_rows:
    if not arabic or not latin:
        continue
    ar2lat_m[norm_ar(arabic)] = latin
    lat2ar_m[norm_lat(latin)] = arabic


def district_key_en(latin: str):
    k = norm_lat(latin)
    return latin if k in lat2ar_d else None


def district_key_ar(arabic: str):
    c = clean_ar_district(arabic)
    # Explicit aliases for Arabic district names missing from DistrictTranslation
    # (keys are clean_ar_district() output).
    if c == norm_ar("صيدا"):  # قرى صيدا — the villages of Saida
        return "Saida"
    if c in ar2lat_d:
        return ar2lat_d[c]
    best, bs = None, 0.0
    for known in KNOWN_AR_D:
        s = SequenceMatcher(None, c, known).ratio()
        if s > bs:
            best, bs = known, s
    return ar2lat_d[best] if bs >= 0.8 else None


def mohafaza_key_en(latin: str):
    return latin if norm_lat(latin) in lat2ar_m else None


def mohafaza_key_ar(arabic: str):
    c = norm_ar(arabic)
    return ar2lat_m.get(c)


# --- matched pairs ----------------------------------------------------------
# (multi-arabic cells are already expanded to one pair per row by match_v6;
#  split defensively anyway and tag subdivision groups for the page)
rows = []
for r in list(wb["English (Full)"].iter_rows(values_only=True))[1:]:
    en, ar, d, m, src = r[0], r[1], r[2], r[3], r[4]
    if not en or not ar:
        continue
    parts = split_ar_cell(ar)
    for a in parts:
        rows.append((en, a, d, m, src))
ar_all_by_group = defaultdict(list)
for en, a, d, m, _src in rows:
    if a not in ar_all_by_group[(en, d, m)]:
        ar_all_by_group[(en, d, m)].append(a)
matched = []
for en, a, d, m, src in rows:
    group = f"{en}|{d}|{m}"
    all_ar = ar_all_by_group[(en, d, m)]
    matched.append({
        "en": en, "ar": a, "ar_all": all_ar,
        "subdiv": len(all_ar) > 1, "subdiv_group": group,
        "district": d, "mohafaza": m,
        "district_ar": lat2ar_d.get(norm_lat(d or "")),
        "mohafaza_ar": lat2ar_m.get(norm_lat(m or "")),
        "source": src,
    })

# one-to-one guard: (normalized arabic, district key, mohafaza key) already matched
matched_triples = set()
for m in matched:
    dk, mk = district_key_en(m["district"]), mohafaza_key_en(m["mohafaza"])
    if dk and mk:
        matched_triples.add((norm_ar(m["ar"]), dk, mk))

# --- review candidates ------------------------------------------------------
review = []
for r in list(wb["Review"].iter_rows(values_only=True))[1:]:
    en, d, m = r[0], r[1], r[2]
    if not en:
        continue
    cands = []
    for cell in r[3:6]:
        if not cell:
            continue
        mm = re.match(r"^(.*?)\s*\[(\d+(?:\.\d+)?)\]\s*$", str(cell).strip())
        if not mm:
            continue
        name, score = mm.group(1).strip(), float(mm.group(2))
        # candidate district hint is inside parens of the arabic name
        hint = re.search(r"\(([^)]+)\)", name)
        cands.append({
            "ar": re.sub(r"\s*\([^)]*\)\s*", "", name).strip(),
            "ar_district": hint.group(1).strip() if hint else None,
            "score": score,
        })
    review.append({
        "en": en, "district": d, "mohafaza": m,
        "district_ar": lat2ar_d.get(norm_lat(d or "")),
        "candidates": cands,
    })

# --- unmatched --------------------------------------------------------------
unmatched_en, unmatched_ar = [], []
for r in list(wb["Unmatched"].iter_rows(values_only=True))[1:]:
    side, name, d, m = r[0], r[1], r[2], r[3]
    if not name:
        continue
    if side == "English":
        unmatched_en.append({
            "en": name, "district": d, "mohafaza": m,
            "dkey": district_key_en(d), "mkey": mohafaza_key_en(m),
        })
    else:
        dk, mk = district_key_ar(d), mohafaza_key_ar(m)
        if dk and mk and (norm_ar(name), dk, mk) in matched_triples:
            continue  # one-to-one: already matched, not unmatched
        unmatched_ar.append({
            "ar": name, "district": d, "mohafaza": m,
            "dkey": dk, "mkey": mk,
        })

# --- district list for filters ----------------------------------------------
districts = []
seen = set()
for latin, arabic, _moh in dist_rows:
    if latin and latin not in seen:
        seen.add(latin)
        districts.append({"key": latin, "label_en": latin,
                          "label_ar": arabic})
districts.sort(key=lambda x: x["label_en"])

data = {
    "generated_from": Path(_args.input).name,
    "counts": {
        "matched": len(matched),
        "review": len(review),
        "unmatched_en": len(unmatched_en),
        "unmatched_ar": len(unmatched_ar),
    },
    "matched": matched,
    "review": review,
    "unmatched_en": unmatched_en,
    "unmatched_ar": unmatched_ar,
    "districts": districts,
}

OUT.parent.mkdir(exist_ok=True)
OUT.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")),
               encoding="utf-8")
print("wrote", OUT, f"{OUT.stat().st_size/1024:.0f} KB")
print("counts:", data["counts"])
print("unmatched_en missing dkey:", sum(1 for u in unmatched_en if not u["dkey"]))
print("unmatched_ar missing dkey:", sum(1 for u in unmatched_ar if not u["dkey"]))
print("unmatched_ar missing mkey:", sum(1 for u in unmatched_ar if not u["mkey"]))
